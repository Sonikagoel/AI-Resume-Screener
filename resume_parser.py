import pymupdf
from docx import Document


def extract_text_from_pdf(pdf_path):
    text = ""

    with pymupdf.open(pdf_path) as pdf:
        for page in pdf:
            text += page.get_text()

    return text
 

def extract_text_from_docx(docx_path):
    document = Document(docx_path)

    text = ""

    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"

    return text


def extract_text_from_uploaded_file(uploaded_file):
    if uploaded_file.name.lower().endswith(".pdf"):
        with pymupdf.open(
            stream=uploaded_file.getvalue(),
            filetype="pdf"
        ) as pdf:
            return "\n".join(page.get_text() for page in pdf)

    elif uploaded_file.name.lower().endswith(".docx"):
        from io import BytesIO

        document = Document(
            BytesIO(uploaded_file.getvalue())
        )

        return "\n".join(
            paragraph.text for paragraph in document.paragraphs
        )

    return ""
